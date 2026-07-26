from docx import Document
import sys

doc = Document(sys.argv[1])
for i, p in enumerate(doc.paragraphs):
    text = p.text.replace("\n", "\\n")
    if text:
        print(f"P{i:03d}\t{p.style.name}\t{text}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti} {len(table.rows)}x{len(table.columns)}")
    for ri, row in enumerate(table.rows):
        vals = [cell.text.replace("\n", " / ") for cell in row.cells]
        print(f"  R{ri}: " + " || ".join(vals))
